# -*- coding: utf-8 -*-

import os
import sys
import io
import asyncio
import traceback
import openpyxl
from datetime import datetime, timedelta
from typing import List, Optional

from rich.console import Console

from selenium import webdriver
from selenium.common.exceptions import (
    StaleElementReferenceException,
    TimeoutException,
)
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.remote.webelement import WebElement
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

from webdriver_manager.chrome import ChromeDriverManager

# ============================================================
# AJUSTE DO PYTHON PATH
# ============================================================

sys.path.append(
    os.path.abspath(
        os.path.join(
            os.path.dirname(__file__),
            "..",
            "..",
            "..",
        )
    )
)


# ============================================================
# IMPORTAÇÕES DO WORKER
# ============================================================

from worker_automate_hub.models.dto.rpa_historico_request_dto import (
    RpaHistoricoStatusEnum,
    RpaRetornoProcessoDTO,
    RpaTagDTO,
    RpaTagEnum,
)

from worker_automate_hub.models.dto.rpa_processo_entrada_dto import (
    RpaProcessoEntradaDTO,
)

from worker_automate_hub.utils.credentials_manager import (
    CredentialsManager,
)

from worker_automate_hub.utils.util import worker_sleep

from worker_automate_hub.api.client import (
    send_file,
    baixa_boleto_antigo_emsys,
)

console = Console()


# ============================================================
# PROCESSO PRINCIPAL
# ============================================================


class BaixarFaturasPortal:

    # ========================================================
    # CONFIGURAÇÕES
    # ========================================================

    CLIENT_SAP = "800"

    TIMEOUT_CURTO = 10
    TIMEOUT_PADRAO = 30
    TIMEOUT_LONGO = 60

    # Tempo máximo para aguardar o Excel
    TIMEOUT_DOWNLOAD = 180

    # ========================================================
    # URL SAP
    # ========================================================

    URL_CUSTOMER_LINE_ITEMS = (
        "https://vhsrkps4ci.sap.simrede.com.br:44300/"
        "sap/bc/ui2/flp?"
        "appState=lean&"
        "sap-client=800&"
        "sap-language=PT"
        "#Customer-manageLineItems"
    )

    # ========================================================
    # LOGIN SAP
    # ========================================================

    X_CLIENT = "//input[@id='CLIENT_FIELD-inner']"

    X_LOGIN_INPUTS = "//input[contains(@class,'loginInputField')]"

    X_LOGIN_BUTTON = (
        "//*[contains(@class,'LOGIN_SUBMIT_BLOCK') "
        "or @id='LOGIN_LINK' "
        "or @id='LOGIN_SUBMIT_BLOCK']"
    )

    # ========================================================
    # POPUP SAP
    # ========================================================

    X_FECHAR_POPUP = "//button[@title='Fechar' " "or @aria-label='Fechar']"

    # ========================================================
    # FILTRO EMPRESA
    # ========================================================

    X_EMPRESA = "//input[contains(@id,'CompanyCode-inner')]"

    # ========================================================
    # FILTRO TIPO DE DOCUMENTO
    # ========================================================

    X_TIPO_DOCUMENTO = "//input[contains(@id,'AccountingDocumentType-inner')]"

    # ========================================================
    # FILTRO DATA DE VENCIMENTO
    # ========================================================

    X_DATA_VENCIMENTO = "//input[contains(@id,'NetDueDate-inner')]"

    X_DATA_VALOR_POPUP = "//input[@placeholder='Valor']"

    X_BOTAO_OK_DATA = "//button[contains(@id,'valueHelpDialog-ok')]"

    # ========================================================
    # BOTÃO INICIAR
    # ========================================================

    X_BOTAO_INICIAR = "//button[contains(@id,'SmartFilterBar-btnGo')]"

    # ========================================================
    # RESULTADO SEM DADOS
    # ========================================================

    X_DADOS_NAO_DISPONIVEIS = (
        "//span["
        "contains(@id,'TableItems-noDataMsg') "
        "and normalize-space(.)='Dados não disponíveis'"
        "]"
    )

    # ========================================================
    # EXPORTAR RELATÓRIO
    # ========================================================

    X_BOTAO_EXPORTAR_RELATORIO = "//span[@id='application-Customer-manageLineItems-component---fin.ar.lineitems.display.s1View--fin.ar.lineitems.display.SmartTableItems-btnExcelExport-internalSplitBtn-textButton-img']"

    # ========================================================
    # CONSTRUTOR
    # ========================================================

    def __init__(
        self,
        task: RpaProcessoEntradaDTO,
        base_url: Optional[str] = None,
    ):
        self.task = task

        self.base_url = base_url.strip() if base_url else self.URL_CUSTOMER_LINE_ITEMS

        # ====================================================
        # HISTÓRICO
        # ====================================================

        self.historico_id = str(
            getattr(
                task,
                "historico_id",
                "",
            )
            or ""
        ).strip()

        # ====================================================
        # CREDENCIAIS SAP
        # ====================================================

        self.user = CredentialsManager().get_by_key("SAP_USER_DRC")

        self.password = CredentialsManager().get_by_key("SAP_PASSWORD_DRC")

        # ====================================================
        # PASTA DE DOWNLOAD
        # ====================================================

        self.pasta_download = os.path.abspath(
            os.path.join(
                os.getcwd(),
                "downloads_faturas_portal",
            )
        )

        os.makedirs(
            self.pasta_download,
            exist_ok=True,
        )

        # ====================================================
        # DRIVER
        # ====================================================

        self.driver: Optional[webdriver.Chrome] = None

        console.print("[INIT] BaixarFaturasPortal inicializado.")

        console.print(f"[INIT] URL SAP: {self.base_url}")

        console.print(f"[INIT] Pasta download: {self.pasta_download}")

        console.print("[INIT] Usuário carregado? " f"{'SIM' if self.user else 'NÃO'}")

        console.print(
            "[INIT] Histórico carregado? " f"{'SIM' if self.historico_id else 'NÃO'}"
        )

    # ========================================================
    # RETORNOS PADRONIZADOS
    # ========================================================

    def retorno_sucesso(
        self,
        mensagem: str,
    ) -> RpaRetornoProcessoDTO:

        console.print(f"[SUCESSO] {mensagem}")

        return RpaRetornoProcessoDTO(
            sucesso=True,
            retorno=mensagem,
            status=RpaHistoricoStatusEnum.Sucesso,
        )

    def retorno_erro(
        self,
        mensagem: str,
    ) -> RpaRetornoProcessoDTO:

        console.print(f"[ERRO_TECNICO] {mensagem}")

        return RpaRetornoProcessoDTO(
            sucesso=False,
            retorno=mensagem,
            status=RpaHistoricoStatusEnum.Falha,
            tags=[RpaTagDTO(descricao=RpaTagEnum.Tecnico)],
        )

    # ========================================================
    # VALIDAÇÕES
    # ========================================================

    def validar_config(self) -> None:

        if not self.user:
            raise RuntimeError("Credencial SAP_USER_DRC não encontrada.")

        if not self.password:
            raise RuntimeError("Credencial SAP_PASSWORD_DRC não encontrada.")

        if not self.base_url:
            raise RuntimeError("URL do SAP não configurada.")

        if not self.historico_id:
            raise RuntimeError("historico_id não informado no processo.")

    # ========================================================
    # LIMPAR PASTA DE DOWNLOAD
    # ========================================================

    def limpar_pasta_download(self) -> None:

        console.print("[DOWNLOAD] Limpando pasta de download antes de iniciar...")

        os.makedirs(
            self.pasta_download,
            exist_ok=True,
        )

        for nome_arquivo in os.listdir(self.pasta_download):

            caminho_arquivo = os.path.join(
                self.pasta_download,
                nome_arquivo,
            )

            if not os.path.isfile(caminho_arquivo):
                continue

            try:

                os.remove(caminho_arquivo)

                console.print(f"[DOWNLOAD] Arquivo removido: {nome_arquivo}")

            except Exception as e:

                raise RuntimeError(
                    "Não foi possível limpar a pasta de download. "
                    f"Arquivo: {caminho_arquivo}. "
                    f"Erro: {type(e).__name__}: {e}"
                ) from e

        console.print("[DOWNLOAD] Pasta de download limpa com sucesso.")

    # ========================================================
    # FLUXO PRINCIPAL
    # ========================================================

    async def iniciar(
        self,
    ) -> RpaRetornoProcessoDTO:

        step = "INIT"

        try:

            # =================================================
            # VALIDAR CONFIG
            # =================================================

            step = "VALIDAR_CONFIG"

            self.validar_config()

            # =================================================
            # LIMPAR DOWNLOAD
            # =================================================

            step = "LIMPAR_DOWNLOAD"

            self.limpar_pasta_download()

            # =================================================
            # ABRIR CHROME
            # =================================================

            step = "ABRIR_CHROME"

            self.abrir_chrome()

            # =================================================
            # ACESSAR SAP
            # =================================================

            step = "ACESSAR_SAP"

            console.print(f"[{step}] Acessando SAP...")

            if not self.driver:
                raise RuntimeError("ChromeDriver não inicializado.")

            self.driver.get(self.base_url)

            await worker_sleep(3)

            # =================================================
            # LOGIN
            # =================================================

            step = "LOGIN"

            login_ok = await self.login()

            if not login_ok:
                raise RuntimeError("Não foi possível concluir o login no SAP.")

            # =================================================
            # ACESSAR APLICAÇÃO
            # =================================================

            step = "ACESSAR_CUSTOMER_LINE_ITEMS"

            await self.acessar_customer_line_items()

            # =================================================
            # PROCESSAMENTO SAP
            # =================================================

            step = "STEPS_SAP"

            mensagem_final = await self.steps_sap()

            return self.retorno_sucesso(mensagem_final)

        except Exception as e:

            tb = traceback.format_exc()

            mensagem = (
                "Falha durante o processamento no SAP. "
                f"Etapa: {step}. "
                f"Erro: {type(e).__name__}: {e}\n"
                f"{tb}"
            )

            return self.retorno_erro(mensagem)

        finally:

            self.fechar_driver()

    # ========================================================
    # CHROME
    # ========================================================

    def abrir_chrome(self) -> None:

        console.print("[CHROME] Inicializando ChromeDriver...")

        service = Service(ChromeDriverManager().install())

        options = webdriver.ChromeOptions()

        options.add_argument("--lang=pt-BR")

        options.add_argument("--log-level=3")

        options.add_argument("--disable-blink-features=AutomationControlled")

        options.add_argument("--disable-infobars")

        options.add_argument("--disable-notifications")

        options.add_argument("--start-maximized")

        # ====================================================
        # CONFIGURAÇÃO DE DOWNLOAD
        # ====================================================

        prefs = {
            "download.default_directory": self.pasta_download,
            "download.prompt_for_download": False,
            "download.directory_upgrade": True,
            "safebrowsing.enabled": True,
        }

        options.add_experimental_option(
            "prefs",
            prefs,
        )

        options.add_experimental_option(
            "excludeSwitches",
            [
                "enable-automation",
                "enable-logging",
            ],
        )

        options.add_experimental_option(
            "useAutomationExtension",
            False,
        )

        self.driver = webdriver.Chrome(
            service=service,
            options=options,
        )

        self.driver.set_page_load_timeout(self.TIMEOUT_LONGO)

        # ====================================================
        # FORÇAR DOWNLOAD NA PASTA
        # ====================================================

        try:

            self.driver.execute_cdp_cmd(
                "Page.setDownloadBehavior",
                {
                    "behavior": "allow",
                    "downloadPath": self.pasta_download,
                },
            )

        except Exception as e:

            console.print(
                "[CHROME][AVISO] Não foi possível configurar " f"download via CDP: {e}"
            )

        # ====================================================
        # REMOVE navigator.webdriver
        # ====================================================

        try:

            self.driver.execute_cdp_cmd(
                "Page.addScriptToEvaluateOnNewDocument",
                {"source": """
                        Object.defineProperty(
                            navigator,
                            'webdriver',
                            {
                                get: () => undefined
                            }
                        );
                    """},
            )

        except Exception:

            pass

        self.driver.maximize_window()

        console.print("[CHROME] Chrome inicializado.")

    # ========================================================
    # FECHAR DRIVER
    # ========================================================

    def fechar_driver(self) -> None:

        if not self.driver:
            return

        try:

            console.print("[CHROME] Encerrando navegador...")

            self.driver.quit()

        except Exception as e:

            console.print(
                "[CHROME][AVISO] Não foi possível " f"encerrar o navegador: {e}"
            )

        finally:

            self.driver = None

    # ========================================================
    # AGUARDAR ELEMENTO
    # ========================================================

    def aguardar_elemento(
        self,
        xpath: str,
        timeout: Optional[int] = None,
        clicavel: bool = False,
    ) -> WebElement:

        if not self.driver:
            raise RuntimeError("ChromeDriver não inicializado.")

        tempo = timeout or self.TIMEOUT_PADRAO

        if clicavel:

            condicao = EC.element_to_be_clickable(
                (
                    By.XPATH,
                    xpath,
                )
            )

        else:

            condicao = EC.visibility_of_element_located(
                (
                    By.XPATH,
                    xpath,
                )
            )

        return WebDriverWait(
            self.driver,
            tempo,
            ignored_exceptions=(StaleElementReferenceException,),
        ).until(condicao)

    # ========================================================
    # BUSCAR ELEMENTOS
    # ========================================================

    def buscar_elementos(
        self,
        xpath: str,
    ) -> List[WebElement]:

        if not self.driver:
            return []

        try:

            return self.driver.find_elements(
                By.XPATH,
                xpath,
            )

        except Exception:

            return []

    # ========================================================
    # ELEMENTO EXISTE
    # ========================================================

    def elemento_existe(
        self,
        xpath: str,
        timeout: int = 3,
    ) -> bool:

        try:

            self.aguardar_elemento(
                xpath=xpath,
                timeout=timeout,
            )

            return True

        except Exception:

            return False

    # ========================================================
    # CLICAR ELEMENTO
    # ========================================================

    def clicar_elemento(
        self,
        xpath: str,
        timeout: Optional[int] = None,
    ) -> None:

        if not self.driver:
            raise RuntimeError("ChromeDriver não inicializado.")

        elemento = self.aguardar_elemento(
            xpath=xpath,
            timeout=timeout,
            clicavel=True,
        )

        try:

            elemento.click()

        except Exception:

            elemento = WebDriverWait(
                self.driver,
                timeout or self.TIMEOUT_PADRAO,
            ).until(
                EC.presence_of_element_located(
                    (
                        By.XPATH,
                        xpath,
                    )
                )
            )

            self.driver.execute_script(
                "arguments[0].click();",
                elemento,
            )

    # ========================================================
    # ALTERAR CLIENT PARA 800
    # ========================================================

    async def alterar_client_para_800(
        self,
    ) -> None:

        try:

            campo_client = self.aguardar_elemento(
                xpath=self.X_CLIENT,
                timeout=self.TIMEOUT_CURTO,
            )

            valor_atual = str(campo_client.get_attribute("value") or "").strip()

            console.print("[LOGIN] Client atual: " f"{valor_atual or 'não informado'}")

            if valor_atual != self.CLIENT_SAP:

                campo_client.clear()

                campo_client.send_keys(self.CLIENT_SAP)

                console.print("[LOGIN] Client alterado para " f"{self.CLIENT_SAP}.")

        except TimeoutException:

            console.print(
                "[LOGIN] Campo de client não apareceu. "
                "O client já pode estar definido pela URL."
            )

    # ========================================================
    # AGUARDAR INPUTS LOGIN
    # ========================================================

    async def aguardar_inputs_login(
        self,
    ) -> List[WebElement]:

        if not self.driver:
            return []

        try:

            WebDriverWait(
                self.driver,
                self.TIMEOUT_PADRAO,
            ).until(
                lambda driver: len(
                    driver.find_elements(
                        By.XPATH,
                        self.X_LOGIN_INPUTS,
                    )
                )
                >= 2
            )

            return self.driver.find_elements(
                By.XPATH,
                self.X_LOGIN_INPUTS,
            )

        except TimeoutException:

            return []

    # ========================================================
    # AGUARDAR BOTÃO LOGIN
    # ========================================================

    async def aguardar_botao_login(
        self,
    ) -> Optional[WebElement]:

        try:

            return self.aguardar_elemento(
                xpath=self.X_LOGIN_BUTTON,
                timeout=self.TIMEOUT_PADRAO,
                clicavel=True,
            )

        except TimeoutException:

            return None

    # ========================================================
    # LOGIN
    # ========================================================

    async def login(self) -> bool:

        try:

            console.print("[LOGIN] Iniciando login...")

            await self.alterar_client_para_800()

            inputs = await self.aguardar_inputs_login()

            if len(inputs) < 2:

                if (
                    self.driver
                    and "#Customer-manageLineItems" in self.driver.current_url
                ):

                    console.print("[LOGIN] Sessão SAP já autenticada.")

                    return True

                console.print(
                    "[LOGIN][ERRO] Campos de usuário " "e senha não encontrados."
                )

                return False

            campo_usuario = inputs[0]
            campo_senha = inputs[1]

            campo_usuario.clear()

            campo_usuario.send_keys(self.user)

            console.print("[LOGIN] Usuário preenchido.")

            await worker_sleep(0.5)

            campo_senha.clear()

            campo_senha.send_keys(self.password)

            console.print("[LOGIN] Senha preenchida.")

            await worker_sleep(0.5)

            botao_login = await self.aguardar_botao_login()

            if not botao_login:

                console.print("[LOGIN][ERRO] Botão de login não encontrado.")

                return False

            try:

                botao_login.click()

            except Exception:

                if not self.driver:
                    return False

                self.driver.execute_script(
                    "arguments[0].click();",
                    botao_login,
                )

            console.print("[LOGIN] Botão de login clicado.")

            await worker_sleep(5)

            if await self.login_apresentou_erro():
                return False

            console.print("[LOGIN] Login concluído.")

            return True

        except Exception:

            console.print("[LOGIN][ERRO] Erro durante login.")

            console.print(traceback.format_exc())

            return False

    # ========================================================
    # VERIFICAR ERRO DE LOGIN
    # ========================================================

    async def login_apresentou_erro(
        self,
    ) -> bool:

        if not self.driver:
            return True

        xpaths_erro = [
            (
                "//*[contains("
                "translate(., "
                "'ABCDEFGHIJKLMNOPQRSTUVWXYZ', "
                "'abcdefghijklmnopqrstuvwxyz'), "
                "'usuário ou senha')]"
            ),
            (
                "//*[contains("
                "translate(., "
                "'ABCDEFGHIJKLMNOPQRSTUVWXYZ', "
                "'abcdefghijklmnopqrstuvwxyz'), "
                "'logon failed')]"
            ),
            (
                "//*[contains("
                "translate(., "
                "'ABCDEFGHIJKLMNOPQRSTUVWXYZ', "
                "'abcdefghijklmnopqrstuvwxyz'), "
                "'dados de logon')]"
            ),
            (
                "//*[contains(@class,'error') "
                "and string-length(normalize-space(.)) > 0]"
            ),
        ]

        for xpath in xpaths_erro:

            elementos = self.buscar_elementos(xpath)

            for elemento in elementos:

                try:

                    if not elemento.is_displayed():
                        continue

                    mensagem = elemento.text.strip()

                    if mensagem:

                        console.print(f"[LOGIN][ERRO_SAP] {mensagem}")

                        return True

                except StaleElementReferenceException:

                    continue

        return False

    # ========================================================
    # FECHAR POPUPS INICIAIS
    # ========================================================

    async def fechar_popups_iniciais(
        self,
    ) -> None:

        for _ in range(3):

            try:

                if not self.elemento_existe(
                    self.X_FECHAR_POPUP,
                    timeout=2,
                ):
                    return

                self.clicar_elemento(
                    self.X_FECHAR_POPUP,
                    timeout=3,
                )

                console.print("[SAP] Popup inicial fechado.")

                await worker_sleep(1)

            except Exception:

                return

    # ========================================================
    # ACESSAR CUSTOMER LINE ITEMS
    # ========================================================

    async def acessar_customer_line_items(
        self,
    ) -> None:

        if not self.driver:
            raise RuntimeError("ChromeDriver não inicializado.")

        console.print("[SAP] Aguardando carregamento do Fiori...")

        await worker_sleep(3)

        await self.fechar_popups_iniciais()

        if "#Customer-manageLineItems" not in self.driver.current_url:

            console.print("[SAP] Redirecionando para " "Customer-manageLineItems...")

            self.driver.get(self.base_url)

        try:

            WebDriverWait(
                self.driver,
                self.TIMEOUT_LONGO,
            ).until(lambda driver: ("#Customer-manageLineItems" in driver.current_url))

        except TimeoutException as e:

            raise RuntimeError(
                "Não foi possível acessar o aplicativo " "Customer-manageLineItems."
            ) from e

        console.print(
            "[SAP] Aplicativo Customer-manageLineItems " "acessado com sucesso."
        )

        console.print(f"[SAP] URL atual: {self.driver.current_url}")

        await worker_sleep(5)

    # ========================================================
    # FILTRO EMPRESA
    # ========================================================

    async def preencher_empresas(
        self,
    ) -> None:

        if not self.driver:
            raise RuntimeError("ChromeDriver não inicializado.")

        console.print("[SAP][EMPRESA] Iniciando preenchimento...")

        # ====================================================
        # EMPRESA 1000
        # ====================================================

        campo_empresa = self.aguardar_elemento(
            xpath=self.X_EMPRESA,
            timeout=self.TIMEOUT_LONGO,
            clicavel=True,
        )

        campo_empresa.click()

        console.print("[SAP][EMPRESA] Digitando 1000...")

        campo_empresa.send_keys("1000")

        await worker_sleep(2)

        campo_empresa.send_keys(Keys.ENTER)

        console.print("[SAP][EMPRESA] Empresa 1000 confirmada.")

        await worker_sleep(2)

        # ====================================================
        # EMPRESA 9999
        # ====================================================

        campo_empresa = self.aguardar_elemento(
            xpath=self.X_EMPRESA,
            timeout=self.TIMEOUT_PADRAO,
            clicavel=True,
        )

        campo_empresa.click()

        console.print("[SAP][EMPRESA] Digitando 9999...")

        campo_empresa.send_keys("9999")

        await worker_sleep(2)

        campo_empresa.send_keys(Keys.ENTER)

        console.print("[SAP][EMPRESA] Empresa 9999 confirmada.")

        await worker_sleep(2)

        console.print("[SAP][EMPRESA] Empresas 1000 e 9999 preenchidas.")

    # ========================================================
    # FILTRO TIPO DE DOCUMENTO
    # ========================================================

    async def preencher_tipo_documento(
        self,
    ) -> None:

        if not self.driver:
            raise RuntimeError("ChromeDriver não inicializado.")

        console.print("[SAP][TIPO_DOCUMENTO] Digitando UE...")

        campo_tipo = self.aguardar_elemento(
            xpath=self.X_TIPO_DOCUMENTO,
            timeout=self.TIMEOUT_LONGO,
            clicavel=True,
        )

        campo_tipo.click()

        campo_tipo.send_keys("UE")

        await worker_sleep(2)

        campo_tipo.send_keys(Keys.ENTER)

        console.print("[SAP][TIPO_DOCUMENTO] UE confirmado.")

        await worker_sleep(2)

    # ========================================================
    # ABRIR POPUP DATA DE VENCIMENTO
    # ========================================================

    async def abrir_popup_data_vencimento(
        self,
    ) -> None:

        if not self.driver:
            raise RuntimeError("ChromeDriver não inicializado.")

        console.print("[SAP][DATA] Abrindo Data vencimento líq...")

        campo_data = self.aguardar_elemento(
            xpath=self.X_DATA_VENCIMENTO,
            timeout=self.TIMEOUT_LONGO,
            clicavel=False,
        )

        self.driver.execute_script(
            """
            arguments[0].scrollIntoView({
                block: 'center',
                inline: 'center'
            });
            """,
            campo_data,
        )

        await worker_sleep(0.5)

        try:

            campo_data.click()

        except Exception:

            self.driver.execute_script(
                "arguments[0].click();",
                campo_data,
            )

        console.print("[SAP][DATA] Clique realizado no campo.")

        try:

            WebDriverWait(
                self.driver,
                self.TIMEOUT_PADRAO,
            ).until(
                EC.visibility_of_element_located(
                    (
                        By.XPATH,
                        self.X_DATA_VALOR_POPUP,
                    )
                )
            )

        except TimeoutException as e:

            raise RuntimeError(
                "O popup da Data vencimento líq. " "não foi aberto."
            ) from e

        console.print("[SAP][DATA] Popup aberto.")

    # ========================================================
    # FILTRO DATA D-1
    # ========================================================

    async def preencher_data_d_1(
        self,
    ) -> None:

        if not self.driver:
            raise RuntimeError("ChromeDriver não inicializado.")

        data_d_1 = (datetime.now() - timedelta(days=1)).strftime("%d%m%Y")

        console.print(f"[SAP][DATA] D-1 calculado: {data_d_1}")

        await self.abrir_popup_data_vencimento()

        campo_valor = self.aguardar_elemento(
            xpath=self.X_DATA_VALOR_POPUP,
            timeout=self.TIMEOUT_LONGO,
            clicavel=True,
        )

        campo_valor.click()

        campo_valor.send_keys(
            Keys.CONTROL,
            "a",
        )

        campo_valor.send_keys(Keys.DELETE)

        campo_valor.send_keys(data_d_1)

        console.print(f"[SAP][DATA] Data digitada: {data_d_1}")

        await worker_sleep(1)

        campo_valor.send_keys(Keys.ENTER)

        console.print("[SAP][DATA] ENTER enviado.")

        await worker_sleep(2)

        botao_ok = self.aguardar_elemento(
            xpath=self.X_BOTAO_OK_DATA,
            timeout=self.TIMEOUT_PADRAO,
            clicavel=True,
        )

        try:

            botao_ok.click()

        except Exception:

            self.driver.execute_script(
                "arguments[0].click();",
                botao_ok,
            )

        console.print("[SAP][DATA] Botão OK clicado.")

        await worker_sleep(2)

        console.print(f"[SAP][DATA] D-1 configurado: {data_d_1}")

    # ========================================================
    # CLICAR BOTÃO INICIAR
    # ========================================================

    async def clicar_botao_iniciar(
        self,
    ) -> None:

        if not self.driver:
            raise RuntimeError("ChromeDriver não inicializado.")

        console.print("[SAP] Aguardando botão Iniciar...")

        botao_iniciar = self.aguardar_elemento(
            xpath=self.X_BOTAO_INICIAR,
            timeout=self.TIMEOUT_LONGO,
            clicavel=True,
        )

        try:

            botao_iniciar.click()

        except Exception:

            self.driver.execute_script(
                "arguments[0].click();",
                botao_iniciar,
            )

        console.print("[SAP] Botão Iniciar clicado.")

        console.print("[SAP] Aguardando carregamento dos resultados...")

        await worker_sleep(10)

    # ========================================================
    # VERIFICAR SE A CONSULTA RETORNOU SEM DADOS
    # ========================================================

    def dados_nao_disponiveis(
        self,
        timeout: int = 3,
    ) -> bool:

        if not self.driver:
            return False

        try:

            elemento = WebDriverWait(
                self.driver,
                timeout,
                poll_frequency=0.5,
                ignored_exceptions=(StaleElementReferenceException,),
            ).until(
                EC.visibility_of_element_located(
                    (
                        By.XPATH,
                        self.X_DADOS_NAO_DISPONIVEIS,
                    )
                )
            )

            mensagem = str(elemento.text or "").strip()

            if mensagem.casefold() == "dados não disponíveis".casefold():

                console.print("=" * 100)
                console.print(
                    "[SAP][SEM_DADOS] Dados não disponíveis."
                )
                console.print(
                    "[SAP][SEM_DADOS] A consulta não retornou registros. "
                    "O fluxo será encerrado como sucesso."
                )
                console.print("=" * 100)

                return True

        except TimeoutException:

            return False

        except StaleElementReferenceException:

            return False

        except Exception as e:

            console.print(
                "[SAP][SEM_DADOS][AVISO] Não foi possível validar "
                f"a mensagem de ausência de dados: {type(e).__name__}: {e}"
            )

            return False

        return False

    # ========================================================
    # AGUARDAR DOWNLOAD
    # ========================================================

    async def aguardar_download_excel(
        self,
    ) -> str:

        console.print("[DOWNLOAD] Aguardando download do Excel...")

        inicio = datetime.now()

        while True:

            # =================================================
            # VERIFICA TIMEOUT
            # =================================================

            tempo_decorrido = (datetime.now() - inicio).total_seconds()

            if tempo_decorrido > self.TIMEOUT_DOWNLOAD:

                raise RuntimeError(
                    "O arquivo Excel não foi baixado dentro "
                    f"de {self.TIMEOUT_DOWNLOAD} segundos."
                )

            # =================================================
            # VERIFICA DOWNLOAD EM ANDAMENTO
            # =================================================

            arquivos = os.listdir(self.pasta_download)

            downloads_em_andamento = [
                arquivo
                for arquivo in arquivos
                if arquivo.lower().endswith(
                    (
                        ".crdownload",
                        ".tmp",
                        ".part",
                    )
                )
            ]

            if downloads_em_andamento:

                await worker_sleep(1)

                continue

            # =================================================
            # PROCURA XLSX
            # =================================================

            arquivos_excel = []

            for arquivo in arquivos:

                caminho = os.path.join(
                    self.pasta_download,
                    arquivo,
                )

                if os.path.isfile(caminho) and arquivo.lower().endswith(".xlsx"):

                    arquivos_excel.append(caminho)

            if arquivos_excel:

                # Pega o mais recente
                caminho_excel = max(
                    arquivos_excel,
                    key=os.path.getmtime,
                )

                # =================================================
                # CONFIRMA QUE O ARQUIVO NÃO ESTÁ MAIS CRESCENDO
                # =================================================

                tamanho_1 = os.path.getsize(caminho_excel)

                await worker_sleep(2)

                tamanho_2 = os.path.getsize(caminho_excel)

                if tamanho_1 > 0 and tamanho_1 == tamanho_2:

                    console.print("[DOWNLOAD] Excel baixado com sucesso:")

                    console.print(f"[DOWNLOAD] {caminho_excel}")

                    console.print(f"[DOWNLOAD] Tamanho: " f"{tamanho_2} bytes")

                    return caminho_excel

            await worker_sleep(1)

    # ========================================================
    # EXPORTAR RELATÓRIO
    # ========================================================

    async def clicar_exportar_relatorio(
        self,
    ) -> str:

        if not self.driver:
            raise RuntimeError("ChromeDriver não inicializado.")

        console.print(
            "[SAP][EXPORTAR] Aguardando botão " "'Exportar para planilha eletrônica'..."
        )

        # ====================================================
        # AGUARDAR BOTÃO
        # ====================================================

        try:

            botao_exportar = WebDriverWait(
                self.driver,
                self.TIMEOUT_LONGO,
                poll_frequency=0.5,
                ignored_exceptions=(StaleElementReferenceException,),
            ).until(
                EC.element_to_be_clickable(
                    (
                        By.XPATH,
                        self.X_BOTAO_EXPORTAR_RELATORIO,
                    )
                )
            )

        except TimeoutException as e:

            raise RuntimeError(
                "O botão 'Exportar para planilha eletrônica' "
                "não ficou disponível após executar a consulta."
            ) from e

        console.print("[SAP][EXPORTAR] Botão encontrado.")

        # ====================================================
        # SCROLL
        # ====================================================

        self.driver.execute_script(
            """
            arguments[0].scrollIntoView({
                block: 'center',
                inline: 'center'
            });
            """,
            botao_exportar,
        )

        await worker_sleep(1)

        # ====================================================
        # CLIQUE NORMAL
        # ====================================================

        try:

            botao_exportar.click()

            console.print("[SAP][EXPORTAR] Clique normal realizado.")

        except Exception as e:

            console.print(
                "[SAP][EXPORTAR] Clique normal falhou: " f"{type(e).__name__}: {e}"
            )

            # =================================================
            # BUSCA NOVAMENTE
            # =================================================

            botao_exportar = WebDriverWait(
                self.driver,
                self.TIMEOUT_PADRAO,
            ).until(
                EC.presence_of_element_located(
                    (
                        By.XPATH,
                        self.X_BOTAO_EXPORTAR_RELATORIO,
                    )
                )
            )

            # =================================================
            # CLIQUE JAVASCRIPT
            # =================================================

            console.print("[SAP][EXPORTAR] Tentando clique via JavaScript...")

            self.driver.execute_script(
                "arguments[0].click();",
                botao_exportar,
            )

            console.print("[SAP][EXPORTAR] Clique via JavaScript realizado.")

        # ====================================================
        # AGUARDAR DOWNLOAD
        # ====================================================

        caminho_excel = await self.aguardar_download_excel()

        return caminho_excel

    # ========================================================
    # NORMALIZAR VALOR DO EXCEL
    # ========================================================

    @staticmethod
    def normalizar_valor_excel(
        valor,
    ) -> str:

        if valor is None:
            return ""

        if isinstance(valor, float):
            if valor != valor:  # NaN
                return ""

            if valor.is_integer():
                return str(int(valor)).strip()

        return str(valor).strip()

    # ========================================================
    # LOCALIZAR CABEÇALHOS NO EXCEL
    # ========================================================

    @staticmethod
    def normalizar_cabecalho(
        valor,
    ) -> str:
        """
        Normaliza cabeçalhos do Excel para permitir pequenas
        variações de espaçamento geradas pelo SAP.

        Exemplos que passam a ser equivalentes:
            "Chave ref. 3"
            "Chave ref.3"
            "Chave  ref.  3"
            "CHAVE REF.3"

        A comparação continua sendo exata após a normalização,
        apenas ignorando diferenças de espaços e caixa.
        """

        if valor is None:
            return ""

        texto = str(valor).strip().casefold()

        # Remove todos os espaços em branco para que, por exemplo:
        # "Chave ref. 3" -> "chaveref.3"
        # "Chave ref.3"  -> "chaveref.3"
        return "".join(texto.split())

    def localizar_colunas_boleto(
        self,
        worksheet,
    ):

        alvo_cliente = self.normalizar_cabecalho("CLIENTE")

        alvo_chave = self.normalizar_cabecalho("Chave ref. 3")

        limite_linhas = min(
            worksheet.max_row,
            50,
        )

        for numero_linha in range(
            1,
            limite_linhas + 1,
        ):

            coluna_cliente = None
            coluna_chave = None

            for numero_coluna in range(
                1,
                worksheet.max_column + 1,
            ):

                valor = worksheet.cell(
                    row=numero_linha,
                    column=numero_coluna,
                ).value

                cabecalho = self.normalizar_cabecalho(valor)

                if cabecalho == alvo_cliente:
                    coluna_cliente = numero_coluna

                elif cabecalho == alvo_chave:
                    coluna_chave = numero_coluna

            if coluna_cliente is not None and coluna_chave is not None:

                return (
                    numero_linha,
                    coluna_cliente,
                    coluna_chave,
                )

        return None

    # ========================================================
    # LER BOLETOS DO EXCEL
    # ========================================================

    def ler_boletos_excel(
        self,
        caminho_excel: str,
    ) -> List[dict]:

        if not os.path.exists(caminho_excel):

            raise RuntimeError(
                "Arquivo Excel não encontrado para leitura dos boletos: "
                f"{caminho_excel}"
            )

        console.print("=" * 100)

        console.print("[BOLETOS] Lendo CLIENTE e Chave ref. 3 do Excel...")

        console.print(f"[BOLETOS] Arquivo: {caminho_excel}")

        workbook = None

        try:

            workbook = openpyxl.load_workbook(
                caminho_excel,
                read_only=True,
                data_only=True,
            )

            worksheet_encontrada = None
            dados_colunas = None

            for worksheet in workbook.worksheets:

                resultado = self.localizar_colunas_boleto(worksheet)

                if resultado:
                    worksheet_encontrada = worksheet
                    dados_colunas = resultado
                    break

            if worksheet_encontrada is None or dados_colunas is None:

                raise RuntimeError(
                    "Não foram encontradas, na mesma aba do Excel, "
                    "as colunas 'CLIENTE' e 'Chave ref. 3'."
                )

            (
                linha_cabecalho,
                coluna_cliente,
                coluna_chave,
            ) = dados_colunas

            console.print("[BOLETOS] Aba encontrada: " f"{worksheet_encontrada.title}")

            console.print("[BOLETOS] Linha do cabeçalho: " f"{linha_cabecalho}")

            console.print("[BOLETOS] Coluna CLIENTE: " f"{coluna_cliente}")

            console.print("[BOLETOS] Coluna Chave ref. 3: " f"{coluna_chave}")

            boletos = []
            total_linhas = 0
            total_sem_cliente = 0
            total_sem_chave = 0
            total_chave_ajustada = 0

            for numero_linha in range(
                linha_cabecalho + 1,
                worksheet_encontrada.max_row + 1,
            ):

                total_linhas += 1

                cliente_raw = worksheet_encontrada.cell(
                    row=numero_linha,
                    column=coluna_cliente,
                ).value

                chave_raw = worksheet_encontrada.cell(
                    row=numero_linha,
                    column=coluna_chave,
                ).value

                cod_bp = self.normalizar_valor_excel(cliente_raw)

                nosso_numero = self.normalizar_valor_excel(chave_raw)

                # Regra solicitada:
                # sem Chave ref. 3, a linha NÃO deve ir para o POST.
                if not nosso_numero:
                    total_sem_chave += 1
                    continue

                # Sem CLIENTE não é possível montar codBp válido.
                if not cod_bp:
                    total_sem_cliente += 1

                    console.print(
                        "[BOLETOS][AVISO] Linha "
                        f"{numero_linha} ignorada: CLIENTE vazio."
                    )

                    continue

                # Regra solicitada:
                # se Chave ref. 3 começar com 000,
                # remover EXATAMENTE os três zeros iniciais.
                if nosso_numero.startswith("000"):

                    nosso_numero = nosso_numero[3:]
                    total_chave_ajustada += 1

                # Se a chave original era somente "000",
                # o resultado ficaria vazio e não deve ser enviado.
                if not nosso_numero:
                    total_sem_chave += 1
                    continue

                boletos.append(
                    {
                        "codBp": cod_bp,
                        "nossoNumero": nosso_numero,
                    }
                )

            console.print("[BOLETOS] Linhas de dados avaliadas: " f"{total_linhas}")

            console.print(
                "[BOLETOS] Linhas ignoradas sem Chave ref. 3: " f"{total_sem_chave}"
            )

            console.print(
                "[BOLETOS] Linhas ignoradas sem CLIENTE: " f"{total_sem_cliente}"
            )

            console.print(
                "[BOLETOS] Chaves iniciadas por 000 e ajustadas: "
                f"{total_chave_ajustada}"
            )

            console.print("[BOLETOS] Boletos válidos para POST: " f"{len(boletos)}")

            return boletos

        except Exception as e:

            raise RuntimeError(
                "Erro ao ler CLIENTE e Chave ref. 3 do Excel. "
                f"Erro: {type(e).__name__}: {e}"
            ) from e

        finally:

            if workbook is not None:
                workbook.close()

    # ========================================================
    # POST - BAIXA BOLETO ANTIGO EMSYS
    # ========================================================

    async def enviar_boletos_baixa_emsys(
        self,
        boletos: List[dict],
    ) -> dict:

        if not boletos:

            console.print(
                "[BOLETOS][AVISO] Nenhum boleto válido encontrado. "
                "O endpoint NÃO será chamado."
            )

            return {
                "enviado": False,
                "quantidade": 0,
                "status_code": None,
                "resposta": None,
            }

        console.print("=" * 100)

        console.print(
            "[BOLETOS][POST] Iniciando chamada para baixa de boleto antigo..."
        )

        console.print(
            "[BOLETOS][POST] A chamada será realizada pelo "
            "worker_automate_hub.api.client."
        )

        console.print(
            "[BOLETOS][POST] Autenticação: Basic Auth centralizada " "no api.client."
        )

        console.print("[BOLETOS][POST] Quantidade de boletos: " f"{len(boletos)}")

        console.print("[BOLETOS][POST] Primeiro item: " f"{boletos[0]}")

        try:

            resposta = await baixa_boleto_antigo_emsys(boletos)

        except Exception as e:

            raise RuntimeError(
                "Falha ao chamar baixa_boleto_antigo_emsys "
                "do worker_automate_hub.api.client. "
                f"Erro: {type(e).__name__}: {e}"
            ) from e

        console.print("[BOLETOS][POST] Chamada concluída com sucesso.")

        if resposta is not None:

            resposta_log = str(resposta)

            limite_log = 2000

            console.print("[BOLETOS][POST] Resposta: " f"{resposta_log[:limite_log]}")

            if len(resposta_log) > limite_log:

                console.print("[BOLETOS][POST][AVISO] " "Resposta truncada no log.")

        else:

            console.print("[BOLETOS][POST] API concluída sem corpo de resposta.")

        return {
            "enviado": True,
            "quantidade": len(boletos),
            "status_code": 200,
            "resposta": resposta,
        }

    # ========================================================
    # ENVIAR ARQUIVO PARA BACKOFFICE
    # ========================================================

    async def enviar_arquivo_backoffice(
        self,
        caminho_arquivo: str,
    ) -> str:

        if not os.path.exists(caminho_arquivo):

            raise RuntimeError(
                "Arquivo Excel não encontrado para envio: " f"{caminho_arquivo}"
            )

        # ====================================================
        # NOME FINAL DO ARQUIVO
        # ====================================================

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        desArquivo = f"faturas_portal_{timestamp}.xlsx"

        console.print("[BACKOFFICE] Preparando arquivo para envio...")

        console.print(f"[BACKOFFICE] Nome: {desArquivo}")

        # ====================================================
        # LER ARQUIVO
        # ====================================================

        with open(
            caminho_arquivo,
            "rb",
        ) as file:

            file_bytes = io.BytesIO(file.read())

        # ====================================================
        # ENVIAR
        # ====================================================

        try:

            console.print("[BACKOFFICE] Enviando arquivo...")

            await send_file(
                self.historico_id,
                desArquivo,
                "xlsx",
                file_bytes,
                file_extension="xlsx",
            )

            console.print("[BACKOFFICE] Arquivo enviado com sucesso.")

        except Exception as e:

            raise RuntimeError(
                "O arquivo de faturas foi gerado com sucesso, "
                "porém ocorreu erro ao enviar para o backoffice. "
                f"Erro: {type(e).__name__}: {e}. "
                "O arquivo foi mantido no dispositivo em: "
                f"{caminho_arquivo}"
            ) from e

        # ====================================================
        # APAGAR SOMENTE DEPOIS DO ENVIO
        # ====================================================

        try:

            os.remove(caminho_arquivo)

            console.print("[BACKOFFICE] Arquivo temporário local removido.")

        except Exception as e:

            console.print(
                "[BACKOFFICE][AVISO] Arquivo enviado, porém "
                "não foi possível apagar o arquivo local: "
                f"{e}"
            )

        return desArquivo

    # ========================================================
    # STEPS SAP
    # ========================================================

    async def steps_sap(
        self,
    ) -> str:

        if not self.driver:
            raise RuntimeError("ChromeDriver não inicializado.")

        console.print("[SAP] Customer-manageLineItems carregado.")

        # ====================================================
        # 1 - EMPRESAS
        # ====================================================

        console.print("[SAP][FILTROS] 1/3 - Empresas...")

        await self.preencher_empresas()

        # ====================================================
        # 2 - TIPO DE DOCUMENTO
        # ====================================================

        console.print("[SAP][FILTROS] 2/3 - Tipo de documento...")

        await self.preencher_tipo_documento()

        # ====================================================
        # 3 - DATA DE VENCIMENTO
        # ====================================================

        console.print("[SAP][FILTROS] 3/3 - Data de vencimento...")

        await self.preencher_data_d_1()

        # ====================================================
        # 4 - INICIAR CONSULTA
        # ====================================================

        console.print("[SAP] Executando consulta...")

        await self.clicar_botao_iniciar()

        # ====================================================
        # 5 - VALIDAR RESULTADO SEM DADOS
        # ====================================================

        if self.dados_nao_disponiveis():

            return "Dados não disponíveis."

        # ====================================================
        # 6 - EXPORTAR RELATÓRIO
        # ====================================================

        console.print("[SAP] Iniciando exportação do relatório...")

        caminho_excel = await self.clicar_exportar_relatorio()

        # ====================================================
        # 7 - LER BOLETOS DO EXCEL
        # ====================================================

        console.print("[SAP] Preparando dados para baixa de boletos no Emsys...")

        boletos = self.ler_boletos_excel(caminho_excel)

        # ====================================================
        # 8 - POST BAIXA BOLETO ANTIGO EMSYS
        # ====================================================

        resultado_post = await self.enviar_boletos_baixa_emsys(boletos)

        # ====================================================
        # 9 - ENVIAR BACKOFFICE
        # ====================================================

        console.print("[SAP] Enviando relatório para o backoffice...")

        nome_enviado = await self.enviar_arquivo_backoffice(caminho_excel)

        # ====================================================
        # FINAL
        # ====================================================

        console.print("[SAP] Fluxo concluído.")

        resposta_endpoint = resultado_post.get("resposta")

        if isinstance(resposta_endpoint, dict):
            resumo_endpoint = " ".join(
                f"{str(chave).replace('_', ' ').capitalize()}: {valor}."
                for chave, valor in resposta_endpoint.items()
            )
        elif resposta_endpoint is not None:
            resumo_endpoint = str(resposta_endpoint)
        else:
            resumo_endpoint = "Sem corpo de resposta."

        return (
            "Consulta realizada com sucesso no SAP. "
            "Empresas: 1000 e 9999. "
            "Tipo de documento: UE. "
            "Data de vencimento: D-1. "
            f"Boletos válidos encontrados: {len(boletos)}. "
            f"POST realizado: {'SIM' if resultado_post['enviado'] else 'NÃO'}. "
            f"Retorno do endpoint: {resumo_endpoint} "
            f"Arquivo enviado para o backoffice: {nome_enviado}"
        )


# ============================================================
# FUNÇÃO CHAMADA PELO WORKER
# ============================================================


async def baixar_faturas_portal(
    task: RpaProcessoEntradaDTO,
) -> RpaRetornoProcessoDTO:

    base_url_sap = (
        "https://vhsrkps4ci.sap.simrede.com.br:44300/"
        "sap/bc/ui2/flp?"
        "appState=lean&"
        "sap-client=800&"
        "sap-language=PT"
        "#Customer-manageLineItems"
    )

    processo = BaixarFaturasPortal(
        task=task,
        base_url=base_url_sap,
    )

    return await processo.iniciar()

