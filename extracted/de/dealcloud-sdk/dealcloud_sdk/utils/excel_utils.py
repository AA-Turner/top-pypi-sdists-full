"""
Excel utilities for exporting and formatting data.

These utilities provide consistent Excel export with formatting,
styling, and multiple sheet support through pandas and openpyxl.
"""

import logging
import re
from pathlib import Path
from typing import TYPE_CHECKING, Any, Callable, Optional, Union

import fsspec
import pandas as pd

if TYPE_CHECKING:
    import polars as pl

logger = logging.getLogger(__name__)

# Constants for Excel compatibility
ILLEGAL_CHARACTERS_RE = r"[\000-\010]|[\013-\014]|[\016-\037]"
INVALID_SHEET_NAME_CHARS = r"[\\*?:/\[\]]"
MAX_SHEET_NAME_LENGTH = 31


def _path_str(filepath: Union[str, Path]) -> str:
    return str(filepath)


def _excel_path_result(filepath: Union[str, Path]) -> Union[str, Path]:
    """Return Path for plain local paths; str for URI destinations."""
    s = _path_str(filepath)
    if "://" in s:
        return s
    return Path(s)


def sanitize_sheet_name(name: str, max_length: int = MAX_SHEET_NAME_LENGTH) -> str:
    """
    Sanitize a string for use as an Excel sheet name.
    
    Removes illegal characters and truncates to max length.
    
    Args:
        name: Proposed sheet name
        max_length: Maximum allowed length (Excel limit is 31)
    
    Returns:
        Sanitized sheet name
    
    Example:
        safe_name = sanitize_sheet_name("Data: Q1/Q2 [2024]")
        # Returns "Data Q1Q2 2024"
    """
    # Remove illegal characters
    sanitized = re.sub(INVALID_SHEET_NAME_CHARS, "", name)
    # Truncate to max length
    return sanitized[:max_length]


def sanitize_cell_value(value: Any) -> Any:
    """
    Sanitize a value for Excel cell compatibility.
    
    Removes illegal control characters that cause Excel errors.
    
    Args:
        value: Cell value to sanitize
    
    Returns:
        Sanitized value
    """
    if isinstance(value, str):
        return re.sub(ILLEGAL_CHARACTERS_RE, "", value)
    return value


def _to_pandas_frame(df: Union[pd.DataFrame, "pl.DataFrame"]) -> pd.DataFrame:
    """Normalize pandas or Polars DataFrame to pandas for openpyxl export."""
    if isinstance(df, pd.DataFrame):
        return df
    try:
        import polars as pl
    except ImportError as e:
        raise ImportError(
            "Polars is not installed. Install with: pip install dealcloud-sdk[polars]"
        ) from e
    if isinstance(df, pl.DataFrame):
        return df.to_pandas()
    raise TypeError(f"Expected pandas or polars DataFrame, got {type(df)!r}")


def sanitize_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """
    Sanitize all string values in a DataFrame for Excel export.
    
    Args:
        df: DataFrame to sanitize
    
    Returns:
        Sanitized DataFrame copy
    
    Example:
        clean_df = sanitize_dataframe(df)
        clean_df.to_excel("output.xlsx")
    """
    return df.replace(ILLEGAL_CHARACTERS_RE, "", regex=True)


def export_to_excel(
    data: Union[pd.DataFrame, "pl.DataFrame", dict[str, pd.DataFrame], dict[str, "pl.DataFrame"]],
    filepath: Union[str, Path],
    sanitize: bool = True,
    index: bool = False,
    **kwargs,
) -> Union[str, Path]:
    """
    Export DataFrame(s) to Excel file.
    
    Supports single DataFrame or dict of DataFrames for multiple sheets.
    Polars DataFrames are converted via ``to_pandas()`` at export time.
    
    Args:
        data: DataFrame or dict of {sheet_name: DataFrame} (pandas or Polars)
        filepath: Output file path
        sanitize: If True, sanitize data for Excel compatibility
        index: If True, include DataFrame index
        **kwargs: Additional arguments passed to to_excel
    
    Returns:
        Path for local targets; URI string for remote targets (e.g. ``s3://``).

    Example:
        # Single sheet
        export_to_excel(df, "output.xlsx")
        
        # Multiple sheets
        export_to_excel({
            "Summary": summary_df,
            "Details": details_df,
        }, "report.xlsx")
    """
    path_str = _path_str(filepath)

    with fsspec.open(path_str, "wb") as out:
        with pd.ExcelWriter(out, engine='openpyxl') as writer:
            if isinstance(data, dict):
                # Dict of DataFrames
                for sheet_name, df in data.items():
                    safe_name = sanitize_sheet_name(sheet_name)
                    pdf = _to_pandas_frame(df)
                    clean_df = sanitize_dataframe(pdf) if sanitize else pdf
                    clean_df.to_excel(writer, sheet_name=safe_name, index=index, **kwargs)
            else:
                # Single DataFrame
                pdf = _to_pandas_frame(data)
                df = sanitize_dataframe(pdf) if sanitize else pdf
                df.to_excel(writer, sheet_name="Sheet1", index=index, **kwargs)

    logger.info(f"Exported data to {path_str}")
    return _excel_path_result(filepath)


def export_with_formatting(
    data: Union[pd.DataFrame, "pl.DataFrame"],
    filepath: Union[str, Path],
    column_widths: Optional[dict[str, int]] = None,
    freeze_panes: Optional[tuple[int, int]] = (1, 0),
    header_style: Optional[dict] = None,
    sanitize: bool = True,
    index: bool = False,
) -> Union[str, Path]:
    """
    Export DataFrame to Excel with formatting.
    
    Args:
        data: DataFrame to export
        filepath: Output file path
        column_widths: Dict of {column_name: width} or {column_index: width}
        freeze_panes: Tuple (row, col) to freeze (default: freeze header row)
        header_style: Dict of style options for header row
        sanitize: If True, sanitize data for Excel compatibility
        index: If True, include DataFrame index
    
    Returns:
        Path for local targets; URI string for remote targets.

    Example:
        export_with_formatting(
            df,
            "output.xlsx",
            column_widths={"Name": 30, "Description": 50},
            freeze_panes=(1, 0),  # Freeze first row
        )
    """
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    
    path_str = _path_str(filepath)
    pdf = _to_pandas_frame(data)
    df = sanitize_dataframe(pdf) if sanitize else pdf

    with fsspec.open(path_str, "wb") as out:
        with pd.ExcelWriter(out, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name="Sheet1", index=index)
            ws = writer.sheets["Sheet1"]
        
            # Apply freeze panes
            if freeze_panes:
                ws.freeze_panes = ws.cell(row=freeze_panes[0] + 1, column=freeze_panes[1] + 1)

            # Apply column widths
            if column_widths:
                for col_ref, width in column_widths.items():
                    if isinstance(col_ref, str):
                        # Column name - find index
                        if col_ref in df.columns:
                            col_idx = df.columns.get_loc(col_ref) + 1
                            if index:
                                col_idx += 1
                            ws.column_dimensions[get_column_letter(col_idx)].width = width
                    else:
                        # Column index
                        ws.column_dimensions[get_column_letter(col_ref + 1)].width = width

            # Apply header style
            if header_style:
                for col in range(1, ws.max_column + 1):
                    cell = ws.cell(row=1, column=col)
                    if 'bold' in header_style:
                        cell.font = Font(bold=header_style['bold'])
                    if 'fill_color' in header_style:
                        cell.fill = PatternFill(
                            start_color=header_style['fill_color'],
                            end_color=header_style['fill_color'],
                            fill_type="solid"
                        )
                    if 'alignment' in header_style:
                        cell.alignment = Alignment(**header_style['alignment'])

    logger.info(f"Exported formatted data to {path_str}")
    return _excel_path_result(filepath)


def create_data_validation_dropdown(
    filepath: Union[str, Path],
    sheet_name: str,
    column: Union[str, int],
    options: list[str],
    start_row: int = 2,
    end_row: Optional[int] = None,
) -> None:
    """
    Add data validation dropdown to an Excel column.
    
    Args:
        filepath: Path to existing Excel file
        sheet_name: Sheet to modify
        column: Column name or letter (e.g., "A" or "Status")
        options: List of valid options
        start_row: First row for validation (default: 2, after header)
        end_row: Last row for validation (None = all rows)
    
    Example:
        create_data_validation_dropdown(
            "output.xlsx",
            "Sheet1",
            "Status",
            ["Active", "Inactive", "Pending"]
        )
    """
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    path_str = _path_str(filepath)
    with fsspec.open(path_str, "rb") as rf:
        wb = load_workbook(rf)
    ws = wb[sheet_name]

    # Get column letter
    if isinstance(column, int):
        col_letter = get_column_letter(column)
    elif len(column) <= 2 and column.isalpha():
        col_letter = column.upper()
    else:
        # Column name - find it
        for col_idx, cell in enumerate(ws[1], 1):
            if cell.value == column:
                col_letter = get_column_letter(col_idx)
                break
        else:
            raise ValueError(f"Column '{column}' not found in sheet")
    
    # Create validation
    if end_row is None:
        end_row = ws.max_row
    
    dv = DataValidation(
        type="list",
        formula1=f'"{",".join(options)}"',
        allow_blank=True,
    )
    dv.error = "Please select from the dropdown options"
    dv.prompt = "Select an option"
    
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}{start_row}:{col_letter}{end_row}")

    with fsspec.open(path_str, "wb") as wf:
        wb.save(wf)
    logger.info(f"Added dropdown validation to column {col_letter} in {path_str}")


def read_excel_sheets(
    filepath: Union[str, Path],
) -> dict[str, pd.DataFrame]:
    """
    Read all sheets from an Excel file.
    
    Args:
        filepath: Path to Excel file
    
    Returns:
        Dict of {sheet_name: DataFrame}
    
    Example:
        sheets = read_excel_sheets("data.xlsx")
        for name, df in sheets.items():
            print(f"Sheet {name}: {len(df)} rows")
    """
    path_str = _path_str(filepath)
    with fsspec.open(path_str, "rb") as f:
        return pd.read_excel(f, sheet_name=None, engine='openpyxl')


def append_to_excel(
    data: Union[pd.DataFrame, "pl.DataFrame"],
    filepath: Union[str, Path],
    sheet_name: str = "Sheet1",
    sanitize: bool = True,
) -> Union[str, Path]:
    """
    Append DataFrame rows to an existing Excel file.
    
    Args:
        data: DataFrame to append
        filepath: Path to existing Excel file
        sheet_name: Sheet to append to
        sanitize: If True, sanitize data for Excel compatibility
    
    Returns:
        Path for local targets; URI string for remote targets.

    Example:
        append_to_excel(new_rows, "data.xlsx")
    """
    from openpyxl import load_workbook

    path_str = _path_str(filepath)
    pdf = _to_pandas_frame(data)
    df = sanitize_dataframe(pdf) if sanitize else pdf

    with fsspec.open(path_str, "rb") as rf:
        wb = load_workbook(rf)
    ws = wb[sheet_name]

    for _, row in df.iterrows():
        ws.append(list(row))

    with fsspec.open(path_str, "wb") as wf:
        wb.save(wf)
    logger.info(f"Appended {len(df)} rows to {path_str}")
    return _excel_path_result(filepath)


