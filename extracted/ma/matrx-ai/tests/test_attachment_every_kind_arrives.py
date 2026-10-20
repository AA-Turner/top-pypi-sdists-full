"""Every attachment kind arrives whole, or a stored notice says why not.

Real test PB-05 (common-docs operations/real-tests/core-chain) attached a Word document to a
Claude Haiku turn and the wire request carried NO block for it and NO notice: the fallback seam
decided per CONTAINER ("Anthropic accepts documents") while the serializer decides per FORMAT
(``DocumentContent.to_anthropic`` only represents PDF and UTF-8 text), so the .docx fell between
them and ``to_anthropic`` returned ``None`` — dropped with a console line only.

The law these guards pin, for every provider style and every kind a person can attach:
  * after ``preprocess_unsupported_media``, each attachment serializes to EXACTLY ONE block
    whose body carries the attachment's fact, or
  * the turn stops with ``MediaFallbackResolutionError`` and a notice is STORED on the request
    snapshot metadata (``media_notices``) and streamed.
A conversion (Office → markdown, PDF → text) also stores its notice.
"""

from __future__ import annotations

import base64
import io
import json
from types import SimpleNamespace

import pytest

from matrx_ai.config import MessageList, TextContent, UnifiedMessage
from matrx_ai.config.media_config import DocumentContent, ImageContent
from matrx_ai.processing.media_fallback import (
    MediaFallbackResolutionError,
    preprocess_unsupported_media,
)
from matrx_ai.providers.provider_io_capabilities import endpoint_serializer_style

FACT = "Clerk code RC-618"
CSV_ROWS = "\n".join(
    ["arrival,shipper,trailer"]
    + [f"{i},Shipper {i},TR-4{i:04d}" for i in range(1, 7)]
    + ["7,Lodi Orchards Cooperative,TR-40917"]
    + [f"{i},Shipper {i},TR-4{i:04d}" for i in range(8, 13)]
)


def _b64(raw: bytes) -> str:
    return base64.b64encode(raw).decode("ascii")


def _pdf() -> bytes:
    import pymupdf

    doc = pymupdf.open()
    for n in range(1, 4):
        page = doc.new_page()
        page.insert_text((72, 72), f"Page {n}." + (f" {FACT}" if n == 3 else ""))
    return doc.tobytes()


def _docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph("Receiving shift handover.")
    d.add_paragraph(f"{FACT} this week.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["field", "value"])
    ws.append(["clerk", FACT])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx() -> bytes:
    import pptx

    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Handover"
    slide.placeholders[1].text = FACT
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _image(fmt: str) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (8, 8), (200, 10, 10)).save(buf, format=fmt)
    return buf.getvalue()


DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

# (label, factory, expects the fact in the serialized body)
KINDS: list[tuple[str, object, bool]] = [
    ("pdf", lambda: DocumentContent(base64_data=_b64(_pdf()), mime_type="application/pdf"), False),
    ("docx", lambda: DocumentContent(base64_data=_b64(_docx()), mime_type=DOCX), True),
    ("xlsx", lambda: DocumentContent(base64_data=_b64(_xlsx()), mime_type=XLSX), True),
    ("pptx", lambda: DocumentContent(base64_data=_b64(_pptx()), mime_type=PPTX), True),
    ("csv", lambda: DocumentContent(base64_data=_b64(CSV_ROWS.encode()), mime_type="text/csv"), True),
    ("txt", lambda: DocumentContent(base64_data=_b64(FACT.encode()), mime_type="text/plain"), True),
    ("md", lambda: DocumentContent(base64_data=_b64(f"# {FACT}".encode()), mime_type="text/markdown"), True),
    ("json", lambda: DocumentContent(base64_data=_b64(json.dumps({"c": FACT}).encode()), mime_type="application/json"), True),
    ("jpeg", lambda: ImageContent(base64_data=_b64(_image("JPEG")), mime_type="image/jpeg"), False),
    ("png", lambda: ImageContent(base64_data=_b64(_image("PNG")), mime_type="image/png"), False),
    ("webp", lambda: ImageContent(base64_data=_b64(_image("WEBP")), mime_type="image/webp"), False),
    ("gif", lambda: ImageContent(base64_data=_b64(_image("GIF")), mime_type="image/gif"), False),
]

ENDPOINTS = ["anthropic_chat", "openai_chat", "google_chat"]


def _serialize(content, endpoint: str):
    style = endpoint_serializer_style(endpoint)
    if isinstance(content, TextContent):
        return {"type": "text", "text": content.text}
    if style == "anthropic":
        return content.to_anthropic()
    if style == "openai_responses":
        return content.to_openai()
    if style == "google":
        return content.to_google()
    method = getattr(content, "to_openai_chat", None)
    return method() if method else None


def _readable_body(block: dict) -> str:
    """The block as text, with any inline base64 payload decoded (Gemini inlineData)."""
    inline = block.get("inlineData") if isinstance(block, dict) else None
    if inline:
        return base64.b64decode(inline["data"]).decode("utf-8", "replace")
    return json.dumps(block)


@pytest.fixture
def ambient(monkeypatch):
    """An AppContext stand-in: a capturing emitter and the metadata dict the snapshot reads."""
    import matrx_connect
    from matrx_connect.context import app_context

    events: list = []

    class _Emitter:
        async def send_data(self, payload):
            events.append(payload)

    ctx = SimpleNamespace(emitter=_Emitter(), metadata={}, request_id="req-1")
    monkeypatch.setattr(matrx_connect, "try_get_app_context", lambda: ctx)
    monkeypatch.setattr(app_context, "try_get_app_context", lambda: ctx)
    ctx.events = events
    return ctx


def _stored_notices(ctx) -> list[dict]:
    staged = ctx.metadata.get("request_snapshot_metadata") or {}
    return list(staged.get("media_notices") or [])


@pytest.mark.parametrize("endpoint", ENDPOINTS)
@pytest.mark.parametrize("label,factory,fact_in_body", KINDS, ids=[k[0] for k in KINDS])
async def test_every_kind_is_one_block_or_one_stored_notice(ambient, endpoint, label, factory, fact_in_body):
    item = factory()
    ml = MessageList(_messages=[UnifiedMessage(role="user", content=[item])])
    try:
        out, _ = await preprocess_unsupported_media(
            ml, endpoint, model_name="test-model", model_supports_vision=True
        )
    except MediaFallbackResolutionError:
        assert _stored_notices(ambient), f"{label} on {endpoint}: stopped with no stored notice"
        return
    blocks = [_serialize(c, endpoint) for m in out.to_list() for c in (m.content or [])]
    real = [b for b in blocks if b]
    assert len(real) == 1, f"{label} on {endpoint}: {len(real)} blocks (silent drop when 0)"
    if fact_in_body:
        body = _readable_body(real[0])
        assert ("RC-618" in body) or ("TR-40917" in body), f"{label} on {endpoint}: body lost its fact"
    converted = any(isinstance(c, TextContent) for m in out.to_list() for c in (m.content or []))
    if converted:
        assert _stored_notices(ambient), f"{label} on {endpoint}: converted with no stored notice"


@pytest.mark.parametrize("endpoint", ENDPOINTS)
@pytest.mark.parametrize("label,factory,_f", KINDS, ids=[k[0] for k in KINDS])
def test_fallback_decision_matches_the_serializer(endpoint, label, factory, _f):
    """The seam and the serializer can never disagree: needs-fallback iff the serializer drops it."""
    from matrx_ai.providers.provider_io_capabilities import content_needs_fallback

    item = factory()
    dropped = _serialize(item, endpoint) is None
    needs = content_needs_fallback(item, endpoint, model_supports_vision=True)
    if dropped:
        assert needs, f"{label} on {endpoint}: the serializer drops it and the seam lets it through"
    if endpoint_serializer_style(endpoint) != "google":
        # Anthropic / OpenAI: exact. Gemini's to_google() inlines ANY bytes, but its
        # inline document slot takes PDF and text/* only, so the seam converts Office
        # and application/* text BEFORE a request Gemini would refuse.
        assert needs is dropped, f"{label} on {endpoint}: seam={needs}, serializer drops={dropped}"


async def test_docx_reaches_anthropic_as_markdown_with_its_name(ambient):
    doc = DocumentContent(
        base64_data=_b64(_docx()), mime_type=DOCX, metadata={"file_name": "handover-note.docx"}
    )
    ml = MessageList(_messages=[UnifiedMessage(role="user", content=[doc])])
    out, _ = await preprocess_unsupported_media(
        ml, "anthropic_chat", model_name="claude-haiku-5-5", model_supports_vision=True
    )
    (only,) = [c for m in out.to_list() for c in m.content]
    assert isinstance(only, TextContent) and "RC-618" in only.text
    assert "handover-note.docx" in only.text
    (notice,) = _stored_notices(ambient)
    assert notice["action"] == "extracted" and "handover-note.docx" in notice["user_message"]
    assert ambient.events and ambient.events[0].media_kind == "document"


async def test_pdf_and_csv_stay_native_beside_a_docx_on_anthropic(ambient):
    pdf = DocumentContent(base64_data=_b64(_pdf()), mime_type="application/pdf")
    csv = DocumentContent(base64_data=_b64(CSV_ROWS.encode()), mime_type="text/csv")
    docx_ = DocumentContent(base64_data=_b64(_docx()), mime_type=DOCX)
    ml = MessageList(_messages=[UnifiedMessage(role="user", content=[pdf, csv, docx_])])
    out, _ = await preprocess_unsupported_media(
        ml, "anthropic_chat", model_name="claude-haiku-5-5", model_supports_vision=True
    )
    kinds = [type(c).__name__ for m in out.to_list() for c in m.content]
    assert kinds == ["DocumentContent", "DocumentContent", "TextContent"]
    csv_block = out.to_list()[0].content[1].to_anthropic()
    assert "7,Lodi Orchards Cooperative,TR-40917" in csv_block["source"]["data"]


def test_anthropic_document_block_is_titled_with_the_file_name():
    doc = DocumentContent(
        base64_data=_b64(CSV_ROWS.encode()), mime_type="text/csv", metadata={"file_name": "inbound-oct16.csv"}
    )
    assert doc.to_anthropic()["title"] == "inbound-oct16.csv"


def test_a_serializer_drop_is_stored_never_only_logged(ambient):
    """Layer 2: whatever slips past the seam and serializes to nothing still leaves a stored notice."""
    lost = DocumentContent(file_id="e820ba04-0fdc-4adf-a139-e98601e6339a", mime_type=DOCX)
    message = UnifiedMessage(role="user", content=[TextContent(text="hi"), lost])
    assert len(message.to_anthropic_blocks()) == 1
    (notice,) = _stored_notices(ambient)
    assert notice["action"] == "dropped" and notice["file_id"] == lost.file_id
