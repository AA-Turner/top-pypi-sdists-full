"""Unsupported media is native, converted to text, or terminates explicitly."""

from __future__ import annotations

import base64

import pytest

from matrx_ai.config import MessageList, TextContent, UnifiedMessage
from matrx_ai.config.media_config import (
    DocumentContent,
    ImageContent,
    VideoContent,
    YouTubeVideoContent,
)
from matrx_ai.processing.media_fallback import (
    MediaFallbackResolutionError,
    assert_media_resolvers_registered,
    handler,
    has_unsupported_media,
    preprocess_unsupported_media,
)


def _msg(*content):
    return MessageList(_messages=[UnifiedMessage(role="user", content=list(content))])


def _flatten(ml):
    return [c for m in ml.to_list() for c in (m.content or [])]


def _make_pdf_bytes(text: str) -> bytes | None:
    try:
        import pymupdf  # noqa: F401
    except Exception:
        try:
            import fitz as pymupdf  # type: ignore
        except Exception:
            return None
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    return doc.tobytes()


async def test_no_unsupported_media_is_noop():
    ml = _msg(TextContent(text="hi"), ImageContent(base64_data="x", mime_type="image/png"))
    # Anthropic accepts text + image → nothing to do, returned unchanged.
    out, usage = await preprocess_unsupported_media(ml, "anthropic_chat", model_name="claude")
    assert usage == []
    assert out is ml  # short-circuit: same object


async def test_missing_youtube_resolver_is_an_explicit_terminal_error(monkeypatch):
    monkeypatch.setattr(handler, "_RESOLVERS", {})
    ml = _msg(TextContent(text="summarize"), YouTubeVideoContent(url="https://youtu.be/abc"))
    assert has_unsupported_media(ml, "openai_chat") is True
    with pytest.raises(MediaFallbackResolutionError, match="youtube"):
        await preprocess_unsupported_media(ml, "openai_chat", model_name="gpt-5")


async def test_youtube_kept_on_google():
    ml = _msg(YouTubeVideoContent(url="https://youtu.be/abc"))
    out, _ = await preprocess_unsupported_media(ml, "google_chat", model_name="gemini")
    kinds = [type(c).__name__ for c in _flatten(out)]
    assert "YouTubeVideoContent" in kinds  # Gemini accepts it natively → untouched


async def test_video_kept_on_native_provider_without_resolver(monkeypatch):
    monkeypatch.setattr(handler, "_RESOLVERS", {})
    ml = _msg(VideoContent(base64_data="eA==", mime_type="video/mp4"))
    out, _ = await preprocess_unsupported_media(ml, "google_chat", model_name="gemini")
    assert out is ml


async def test_video_without_conversion_resolver_is_terminal(monkeypatch):
    monkeypatch.setattr(handler, "_RESOLVERS", {})
    ml = _msg(VideoContent(base64_data="eA==", mime_type="video/mp4"))
    with pytest.raises(MediaFallbackResolutionError, match="video"):
        await preprocess_unsupported_media(ml, "openai_chat", model_name="gpt-5")


async def test_resolver_exception_is_terminal(monkeypatch):
    async def broken_resolver(content, ctx):
        raise RuntimeError("conversion exploded")

    monkeypatch.setattr(
        handler,
        "_RESOLVERS",
        {handler.MediaContainer.VIDEO: broken_resolver},
    )
    ml = _msg(VideoContent(base64_data="eA==", mime_type="video/mp4"))
    with pytest.raises(MediaFallbackResolutionError, match="conversion exploded"):
        await preprocess_unsupported_media(ml, "openai_chat", model_name="gpt-5")


def test_required_resolver_assertion_fails_loud(monkeypatch):
    monkeypatch.setattr(handler, "_RESOLVERS", {})
    with pytest.raises(RuntimeError, match="image, video"):
        assert_media_resolvers_registered(
            {handler.MediaContainer.IMAGE, handler.MediaContainer.VIDEO}
        )


async def test_text_only_model_routes_image_through_fallback(monkeypatch):
    async def image_resolver(content, ctx):
        return handler.ResolverResult(
            replacement=TextContent(text="[Image OCR]\nvisible words"),
            action="extracted",
            user_message="Image converted to text.",
        )

    monkeypatch.setattr(
        handler,
        "_RESOLVERS",
        {handler.MediaContainer.IMAGE: image_resolver},
    )
    ml = _msg(ImageContent(base64_data="eA==", mime_type="image/png"))
    out, _ = await preprocess_unsupported_media(
        ml,
        "openai_chat",
        model_name="text-only",
        model_supports_vision=False,
    )
    flat = _flatten(out)
    assert not any(isinstance(item, ImageContent) for item in flat)
    assert any(isinstance(item, TextContent) and "visible words" in item.text for item in flat)


async def test_pdf_extracted_to_text_on_provider_that_cant_take_documents(monkeypatch):
    pdf = _make_pdf_bytes("Hello from a test PDF document")
    if pdf is None:
        pytest.skip("pymupdf not available to build a test PDF")
    import pytesseract

    def fail_if_ocr_runs(*args, **kwargs):
        raise AssertionError("native-text fallback must not require OCR")

    monkeypatch.setattr(pytesseract, "image_to_string", fail_if_ocr_runs)
    doc = DocumentContent(
        base64_data=base64.b64encode(pdf).decode(), mime_type="application/pdf"
    )
    ml = _msg(TextContent(text="read this"), doc)
    # groq is a chat-completions provider — it drops documents; the resolver should
    # extract the text instead.
    out, _ = await preprocess_unsupported_media(ml, "groq_chat", model_name="llama")
    flat = _flatten(out)
    kinds = [type(c).__name__ for c in flat]
    assert "DocumentContent" not in kinds  # the PDF was replaced...
    extracted = [c for c in flat if isinstance(c, TextContent) and "extracted to text" in c.text]
    assert extracted, "PDF should have been extracted into a TextContent"
    assert "Hello from a test PDF" in extracted[0].text


async def test_pdf_kept_on_provider_that_accepts_documents():
    doc = DocumentContent(base64_data="x", mime_type="application/pdf")
    ml = _msg(doc)
    # Anthropic accepts documents natively → untouched (no extraction).
    out, _ = await preprocess_unsupported_media(ml, "anthropic_chat", model_name="claude")
    assert [type(c).__name__ for c in _flatten(out)] == ["DocumentContent"]


def _make_docx_bytes(text: str) -> bytes | None:
    try:
        import io

        import docx
    except Exception:
        return None
    d = docx.Document()
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


async def test_docx_extracted_to_text():
    raw = _make_docx_bytes("Hello from a Word document")
    if raw is None:
        pytest.skip("python-docx not available")
    doc = DocumentContent(
        base64_data=base64.b64encode(raw).decode(),
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    out, _ = await preprocess_unsupported_media(_msg(doc), "groq_chat", model_name="llama")
    flat = _flatten(out)
    assert "DocumentContent" not in [type(c).__name__ for c in flat]
    text = next(c.text for c in flat if isinstance(c, TextContent))
    assert "Hello from a Word document" in text


async def test_xlsx_extracted_to_text():
    # Excel was previously a hard drop ("this document type isn't extractable").
    # The Office codec now reads it to a markdown table for the model.
    import io

    try:
        import openpyxl
    except Exception:
        pytest.skip("openpyxl not available")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Revenue"])
    ws.append(["North", 100])
    buf = io.BytesIO()
    wb.save(buf)
    doc = DocumentContent(
        base64_data=base64.b64encode(buf.getvalue()).decode(),
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    out, _ = await preprocess_unsupported_media(_msg(doc), "groq_chat", model_name="llama")
    flat = _flatten(out)
    assert "DocumentContent" not in [type(c).__name__ for c in flat]
    text = next(c.text for c in flat if isinstance(c, TextContent))
    assert "Region" in text and "North" in text


async def test_plain_text_document_extracted():
    raw = b"plain text body, line one\nline two"
    doc = DocumentContent(base64_data=base64.b64encode(raw).decode(), mime_type="text/plain")
    out, _ = await preprocess_unsupported_media(_msg(doc), "groq_chat", model_name="llama")
    text = next(c.text for c in _flatten(out) if isinstance(c, TextContent))
    assert "plain text body, line one" in text


async def test_unsupported_document_type_is_explicit_terminal_error():
    doc = DocumentContent(base64_data="eA==", mime_type="application/x-unknown-binary")
    with pytest.raises(MediaFallbackResolutionError, match="document"):
        await preprocess_unsupported_media(_msg(doc), "groq_chat", model_name="llama")


async def test_emits_inline_media_notice(monkeypatch):
    captured = []

    class _Emitter:
        async def send_data(self, payload):
            captured.append(payload)

    from types import SimpleNamespace

    import matrx_connect

    monkeypatch.setattr(
        matrx_connect, "try_get_app_context", lambda: SimpleNamespace(emitter=_Emitter())
    )
    ml = _msg(YouTubeVideoContent(url="https://youtu.be/abc"))
    with pytest.raises(MediaFallbackResolutionError):
        await preprocess_unsupported_media(ml, "openai_chat", model_name="gpt-5")
    assert captured, "a media notice should have been emitted"
    notice = captured[0]
    assert notice.media_kind == "youtube"
    assert notice.action == "dropped"
    assert "YouTube" in notice.user_message
